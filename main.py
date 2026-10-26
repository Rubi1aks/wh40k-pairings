import os
import json
import urllib.request
import openpyxl
import tempfile
from fastapi import FastAPI, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel
import minimax
import re

app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/")
def read_root():
    return FileResponse("static/index.html")

class SheetRequest(BaseModel):
    url: str

class GenerateRequest(BaseModel):
    url: str
    sheet_name: str
    team_size: int
    format_type: str
    scale: int

def _extract_gsheet_id(url):
    m = re.search(r'/spreadsheets/d/([a-zA-Z0-9_-]+)', url)
    if not m: raise ValueError("Неверная ссылка")
    return m.group(1)

def download_gsheet(url):
    gid = _extract_gsheet_id(url)
    export_url = f'https://docs.google.com/spreadsheets/d/{gid}/export?format=xlsx'
    req = urllib.request.Request(export_url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=30) as r:
        data = r.read()
    if not data.startswith(b'PK'): raise ValueError("Google вернул не xlsx. Проверьте доступ.")
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    with os.fdopen(fd, 'wb') as f: f.write(data)
    return path

@app.post("/api/sheets")
def get_sheets(req: SheetRequest):
    try:
        path = download_gsheet(req.url)
        wb = openpyxl.load_workbook(path, read_only=True)
        os.remove(path)
        return {"sheets": wb.sheetnames}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/generate")
def generate(req: GenerateRequest):
    try:
        path = download_gsheet(req.url)
        wb = openpyxl.load_workbook(path, data_only=True)
        os.remove(path)
        
        if req.sheet_name not in wb.sheetnames:
            raise ValueError("Лист не найден")
        
        ws = wb[req.sheet_name]
        players, enemies = minimax.validate_table(ws)
        our = [p[0] for p in players]
        our_row = {p[0]: p[1] for p in players}
        enemy_names = [e[0] for e in enemies]
        cM = {e[0]: e[1] for e in enemies}
        cN = {e[0]: e[2] for e in enemies}
        
        sm = {p: {} for p in our}; sn = {p: {} for p in our}
        for p in our:
            for e in enemy_names:
                mv = ws.cell(row=our_row[p], column=cM[e]).value
                nv = ws.cell(row=our_row[p], column=cN[e]).value
                sm[p][e] = minimax._to_score(mv)
                sn[p][e] = minimax._to_score(nv)

        fmt = req.format_type
        if req.team_size in (3, 5): fmt = 'champ_champ'
            
        report = minimax.generate_pairing(our, enemy_names, sm, sn, req.scale, fmt)
        
        try:
            with open('winrates.json', 'r', encoding='utf-8') as f:
                report['winrates'] = json.load(f)
        except: pass
        
        return report
        except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Ошибка при генерации: {str(e)}")
